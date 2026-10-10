#!/usr/bin/env python3
"""Build the layer-comparison Excel workbook from agg_result.json.

Usage: python3 build_xlsx.py <config.py>     (same config as agg_and_verify.py)

Reads cfg.OUTPUT (agg json produced by agg_and_verify.py — all gates already
passed there). Optional cfg.BUILD dict:
  xlsx_out   output path (default: next to OUTPUT, named layer_compare_<a>_vs_<b>.xlsx)
  titles     {run_name: sheet title} overrides
  extra_notes {run_name: [str,...], 'comparison': [str,...]} task-specific notes
Auto notes (scope/provenance/csv-correspondence/window rules) are generated from
the data; extra_notes are appended after them.
"""
import importlib.util
import json
import os
import sys

XLSX_SKILL_DIR = "/Users/yujinqi/.zcode/cli/plugins/cache/zcode-plugins-official/document-skills/0.1.4/skills/xlsx"
sys.path.insert(0, os.path.join(XLSX_SKILL_DIR, "templates"))
from base import *  # noqa: F401,F403
from openpyxl import Workbook
from openpyxl.styles import Alignment

ROOT = os.path.expanduser('~/workspace')


def P(p):
    return p if os.path.isabs(p) else os.path.join(ROOT, p)


cfg_spec = importlib.util.spec_from_file_location('cfg', sys.argv[1])
cfg = importlib.util.module_from_spec(cfg_spec)
cfg_spec.loader.exec_module(cfg)

agg = json.load(open(P(cfg.OUTPUT)))
a_name, b_name = agg['run_names']
build = getattr(cfg, 'BUILD', {})
OUT = P(build.get('xlsx_out') or os.path.join(
    os.path.dirname(P(cfg.OUTPUT)), f'layer_compare_{a_name}_vs_{b_name}.xlsx'))
TITLES = build.get('titles', {})
EXTRA = build.get('extra_notes', {})

MEAN_FIELDS = agg['mean_fields']
DETAIL_HEADERS = {
    'aicore_time_us': 'aicore_time(us)', 'aic_total_cycles': 'aic_total_cycles',
    'aic_mac_time_us': 'aic_mac_time(us)', 'aic_mac_ratio': 'aic_mac_ratio',
    'aic_scalar_time_us': 'aic_scalar_time(us)', 'aic_scalar_ratio': 'aic_scalar_ratio',
    'aic_mte1_time_us': 'aic_mte1_time(us)', 'aic_mte1_ratio': 'aic_mte1_ratio',
    'aic_mte2_time_us': 'aic_mte2_time(us)', 'aic_mte2_ratio': 'aic_mte2_ratio',
    'aic_fixpipe_time_us': 'aic_fixpipe_time(us)', 'aic_fixpipe_ratio': 'aic_fixpipe_ratio',
    'aic_icache_miss_rate': 'aic_icache_miss_rate',
    'aiv_time_us': 'aiv_time(us)', 'aiv_total_cycles': 'aiv_total_cycles',
    'aiv_vec_time_us': 'aiv_vec_time(us)', 'aiv_vec_ratio': 'aiv_vec_ratio',
    'aiv_scalar_time_us': 'aiv_scalar_time(us)', 'aiv_scalar_ratio': 'aiv_scalar_ratio',
    'aiv_mte2_time_us': 'aiv_mte2_time(us)', 'aiv_mte2_ratio': 'aiv_mte2_ratio',
    'aiv_mte3_time_us': 'aiv_mte3_time(us)', 'aiv_mte3_ratio': 'aiv_mte3_ratio',
    'aiv_icache_miss_rate': 'aiv_icache_miss_rate',
    'cube_utilization_pct': 'cube_utilization(%)',
}
FIELD_FMT = {}
for f in MEAN_FIELDS:
    if f.endswith('_ratio') or f.endswith('_rate'):
        FIELD_FMT[f] = '0.0%'
    elif f.endswith('_cycles'):
        FIELD_FMT[f] = '#,##0'
    elif f == 'cube_utilization_pct':
        FIELD_FMT[f] = '0.0'
    else:
        FIELD_FMT[f] = '#,##0.00'

HEADERS = (['模块', '切分策略', '层', '算子名称', '优化思路', 'Start Time(us)',
            'Duration(us)', '个数', '总时间(us)', '占比', 'Wait Time(us)']
           + [DETAIL_HEADERS[f] for f in MEAN_FIELDS[1:]])
LAST_COL = len(HEADERS) + 1
COL_SPEC = {
    7: ('start_us', '@'), 8: ('duration_us', '#,##0.00'), 9: ('count1', '#,##0'),
    10: ('duration_us', '#,##0.00'), 11: ('share', '0.00%'),
}
for j, f in enumerate(MEAN_FIELDS):
    COL_SPEC[12 + j] = (f, FIELD_FMT[f])

MODULE_ORDER = agg['module_order']

wb = Workbook()
wb.properties.creator = "Z.ai"


def wnum(ws, r, c, v, fmt):
    if v is None:
        return
    cell = ws.cell(row=r, column=c, value=v)
    cell.alignment = Alignment(horizontal='right', vertical='center')
    cell.number_format = fmt


def wstart(ws, r, c, v):
    if v is None:
        return
    cell = ws.cell(row=r, column=c, value=str(v))
    cell.alignment = Alignment(horizontal='right', vertical='center')
    cell.number_format = '@'


def style_subtotal_row(ws, row_num, last_col):
    for col in range(2, last_col + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = fill_total()
        cell.font = font_subheader()
    ws.row_dimensions[row_num].height = ROW_HEIGHTS["total"]


def module_tag(rows):
    tags = {r['layer'] for r in rows}
    return tags.pop() if len(tags) == 1 else '—'


def build_run_sheet(ws, run, title):
    data = agg[run]
    setup_sheet(ws, title=title, last_col=LAST_COL)
    for ci, h in enumerate(HEADERS, start=2):
        ws.cell(row=4, column=ci, value=h)
    style_header_row(ws, 4, 2, LAST_COL)

    r = 5
    di = 0
    grand = {'count': 0, 'total': 0.0, 'share': 0.0}
    for mod in MODULE_ORDER:
        rows = [x for x in data['rows'] if x['module'] == mod]
        if not rows:
            continue
        tag = module_tag(rows)
        for x in rows:
            ws.cell(row=r, column=2, value=x['module'])
            ws.cell(row=r, column=3, value=x['shard'])
            ws.cell(row=r, column=4, value=x['layer'])
            ws.cell(row=r, column=5, value=x['name'])
            if x['idea']:
                ws.cell(row=r, column=6, value=x['idea'])
            wstart(ws, r, 7, x['start_us'])
            for c, (field, fmt) in COL_SPEC.items():
                if c == 7:
                    continue
                v = 1 if field == 'count1' else x.get(field)
                wnum(ws, r, c, v, fmt)
            style_data_row(ws, r, 2, LAST_COL, di)
            r += 1
            di += 1
        st = data['subtotals'][mod]
        label = f'{mod}（{tag}）小计' if tag != '—' else f'{mod} 小计'
        ws.cell(row=r, column=2, value=label)
        wnum(ws, r, 9, st['count'], '#,##0')
        wnum(ws, r, 10, st['total_us'], '#,##0.00')
        wnum(ws, r, 11, st['share'], '0.00%')
        style_subtotal_row(ws, r, LAST_COL)
        grand['count'] += st['count']
        grand['total'] += st['total_us']
        grand['share'] += st['share']
        r += 1

    ws.cell(row=r, column=2, value='本表合计（选定层口径）')
    wnum(ws, r, 9, grand['count'], '#,##0')
    wnum(ws, r, 10, grand['total'], '#,##0.00')
    wnum(ws, r, 11, grand['share'], '0.00%')
    style_total_row(ws, r, 2, LAST_COL)
    grand_row = r

    nr = grand_row + 3
    scope_txt = '、'.join(
        f"{m}={'—' if module_tag([x for x in data['rows'] if x['module'] == m]) == '—' else module_tag([x for x in data['rows'] if x['module'] == m])}"
        for m in MODULE_ORDER if any(x['module'] == m for x in data['rows']))
    notes = [
        f"单步稳态数据（{run}）：每模块取一个代表层 + 特殊区段整段（{scope_txt}）；每行 = 该层时间窗"
        f"（组件首算子 start ~ 末算子 end）内一个算子实例，无任何平均。全步共 {data['n_ops_full']} 算子 / "
        f"{data['total_us_full']:.1f}us；占比分母 = 全步算子总时长，本表为选定层口径故 Σ占比 < 1。",
        "kernel_details 对应：Start Time(us) 为文本列、保留 csv 原值精度，可直接在对应 kernel_details.csv "
        "中 grep 对查；agg_and_verify.py 已程序化校验本表全部行的 (Start Time, Duration) 逐行存在于 csv。"
        "通信行为 AivKernel 展开去重后的 hcom 汇总行。",
        "层窗规则：组件匹配算子恒归属本层；窗内未匹配算子（通信/胶水）按窗包含归属。层间接缝存在 ~0.2us "
        "流重叠（下一层首算子可早于本层末算子结束启动），相邻辅流算子次序有 ±1us 级抖动，均为真实数据。",
        f"行序：层窗内严格按 Start Time(us) 升序。层列：D*=draft/MTP 层、L*=target 层、—=非层内区段"
        f"（如模块间交接）。未展开范围：其余各层与层窗外胶水（全步逐实例明细在 agg_result.json）。",
        "切分策略：该模块算子的并行切分（单卡视角）。SP=序列并行交换（allGather/reduceScatter）；"
        "CP=context parallel；EP=专家并行（dispatch/combine 内嵌于 MoeDistribute 算子）。"
        "优化思路列仅每组（模块+切分+算子）首个实例标注，便于筛选。",
    ] + EXTRA.get(run, [])
    for i, t in enumerate(notes):
        c = ws.cell(row=nr + i, column=2, value=t)
        c.font = font_caption()
        c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=False)

    auto_fit_columns(ws, min_width=8, max_width=28, header_row=4, data_start_row=5)
    auto_fit_row_heights(ws, header_row=4, data_start_row=5)
    ws.freeze_panes = 'F5'
    return grand_row


ws1 = wb.active
ws1.title = a_name
g1 = build_run_sheet(ws1, a_name, TITLES.get(
    a_name, f'算子明细（每模块代表层）— {a_name}（逐实例）'))

ws2 = wb.create_sheet(b_name)
g2 = build_run_sheet(ws2, b_name, TITLES.get(
    b_name, f'算子明细（每模块代表层）— {b_name}（逐实例）'))

# ---------- comparison sheet ----------
C_HEADERS = ['模块', '层', '算子名称', f'{a_name} Start(us)', f'{a_name} 耗时(us)',
             f'{b_name} Start(us)', f'{b_name} 耗时(us)', 'Δ(us)', 'Δ幅度/状态']
C_LAST = len(C_HEADERS) + 1
ws3 = wb.create_sheet('算子对比')
setup_sheet(ws3, title=f'算子级逐实例对比 — {a_name} vs {b_name}（每模块代表层，按 算子+层内序 配对）',
            last_col=C_LAST)
for ci, h in enumerate(C_HEADERS, start=2):
    ws3.cell(row=4, column=ci, value=h)
style_header_row(ws3, 4, 2, C_LAST)

r = 5
di = 0
cg = {a_name: 0.0, b_name: 0.0}
for mod in MODULE_ORDER:
    rows = [c for c in agg['comparison'] if c['module'] == mod]
    if not rows:
        continue
    st_a = agg[a_name]['subtotals'][mod]
    st_b = agg[b_name]['subtotals'][mod]
    tag = module_tag([x for x in agg[a_name]['rows'] if x['module'] == mod]
                     or [x for x in agg[b_name]['rows'] if x['module'] == mod])
    for c in rows:
        ws3.cell(row=r, column=2, value=mod)
        ws3.cell(row=r, column=3, value=c['layer'])
        ws3.cell(row=r, column=4, value=c['name'])
        wstart(ws3, r, 5, c[f'{a_name}_start'])
        wnum(ws3, r, 6, c[f'{a_name}_dur'], '#,##0.00')
        wstart(ws3, r, 7, c[f'{b_name}_start'])
        wnum(ws3, r, 8, c[f'{b_name}_dur'], '#,##0.00')
        wnum(ws3, r, 9, c['delta'], '#,##0.00')
        if c['delta_pct'] is not None:
            wnum(ws3, r, 10, c['delta_pct'], '+0.0%;-0.0%')
        else:
            cell = ws3.cell(row=r, column=10,
                            value='新增' if c[f'{a_name}_dur'] is not None else '仅基线')
            cell.alignment = Alignment(horizontal='right', vertical='center')
        style_data_row(ws3, r, 2, C_LAST, di)
        for rn in (a_name, b_name):
            if c[f'{rn}_dur'] is not None:
                cg[rn] += c[f'{rn}_dur']
        r += 1
        di += 1
    label = (f'{mod}（{tag}）小计（{st_a["count"]} vs {st_b["count"]} 行）'
             if tag != '—' else f'{mod} 小计（{st_a["count"]} vs {st_b["count"]} 行）')
    ws3.cell(row=r, column=2, value=label)
    wnum(ws3, r, 6, st_a['total_us'], '#,##0.00')
    wnum(ws3, r, 8, st_b['total_us'], '#,##0.00')
    wnum(ws3, r, 9, st_a['total_us'] - st_b['total_us'], '#,##0.00')
    wnum(ws3, r, 10, (st_a['total_us'] - st_b['total_us']) / st_b['total_us']
         if st_b['total_us'] > 0 else None, '+0.0%;-0.0%')
    style_subtotal_row(ws3, r, C_LAST)
    r += 1

ws3.cell(row=r, column=2, value='本表合计（选定层口径）')
wnum(ws3, r, 6, cg[a_name], '#,##0.00')
wnum(ws3, r, 8, cg[b_name], '#,##0.00')
wnum(ws3, r, 9, cg[a_name] - cg[b_name], '#,##0.00')
wnum(ws3, r, 10, (cg[a_name] - cg[b_name]) / cg[b_name], '+0.0%;-0.0%')
style_total_row(ws3, r, 2, C_LAST)
comp_grand_row = r

nr = comp_grand_row + 3
notes3 = [
    f"配对口径：每行一对实例（无平均），配对键 = 模块+层+算子名+层窗内第 k 次出现；通信按类型+序号配对"
    f"（hcom 任务号跨轮不同，不作配对键）。单边行：新增={a_name} 特有，仅基线={b_name} 特有。",
    f"两侧 Start Time(us) 均为对应 kernel_details.csv 原值（文本列），可分别对查；行序按 {a_name} 侧 "
    f"Start Time（无则 {b_name} 侧）。小计行 = 该模块选定层全部实例 Σ耗时。",
] + EXTRA.get('comparison', [])
for i, t in enumerate(notes3):
    c = ws3.cell(row=nr + i, column=2, value=t)
    c.font = font_caption()
    c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=False)

auto_fit_columns(ws3, min_width=8, max_width=28, header_row=4, data_start_row=5)
auto_fit_row_heights(ws3, header_row=4, data_start_row=5)
ws3.freeze_panes = 'E5'

# ---------- Review sheet ----------
ws4 = wb.create_sheet('Review')
ws4.sheet_properties.tabColor = "FFC000"
setup_sheet(ws4, title='交叉校验（Cross-Validation）', last_col=5)
for ci, h in enumerate(['校验项', '期望值', '实际值（公式）', '状态'], start=2):
    ws4.cell(row=4, column=ci, value=h)
style_header_row(ws4, 4, 2, 5)

tot = {rn: sum(s['total_us'] for s in agg[rn]['subtotals'].values()) for rn in (a_name, b_name)}
cnt = {rn: sum(s['count'] for s in agg[rn]['subtotals'].values()) for rn in (a_name, b_name)}
checks = [
    (f'{a_name} 本表合计 总时间(us)', round(tot[a_name], 3), f"={a_name}!J{g1}"),
    (f'{a_name} 本表行数', cnt[a_name], f"={a_name}!I{g1}"),
    (f'{b_name} 本表合计 总时间(us)', round(tot[b_name], 3), f"={b_name}!J{g2}"),
    (f'{b_name} 本表行数', cnt[b_name], f"={b_name}!I{g2}"),
    (f'对比表 Σ{a_name} = {a_name} 表合计', round(tot[a_name], 3), f"='算子对比'!F{comp_grand_row}"),
    (f'对比表 Σ{b_name} = {b_name} 表合计', round(tot[b_name], 3), f"='算子对比'!H{comp_grand_row}"),
]
r = 5
for i, (name, exp, formula) in enumerate(checks):
    ws4.cell(row=r, column=2, value=name)
    ec = ws4.cell(row=r, column=3, value=exp)
    ec.alignment = Alignment(horizontal='right', vertical='center')
    if isinstance(exp, float):
        ec.number_format = '#,##0.00'
    fc = ws4.cell(row=r, column=4, value=formula)
    fc.alignment = Alignment(horizontal='right', vertical='center')
    fc.number_format = '#,##0.00' if isinstance(exp, float) else '#,##0'
    ws4.cell(row=r, column=5, value=f'=IF(ABS(C{r}-D{r})<0.05,"PASS","FAIL")')
    style_data_row(ws4, r, 2, 5, i)
    ws4.cell(row=r, column=5).alignment = Alignment(horizontal='center', vertical='center')
    r += 1

auto_fit_columns(ws4, min_width=8, max_width=40, header_row=4, data_start_row=5)
auto_fit_row_heights(ws4, header_row=4, data_start_row=5)

wb.save(OUT)
print('saved', OUT)
print('grands:', g1, g2, comp_grand_row)
print(f'shown: {a_name} {tot[a_name]:.1f}us/{cnt[a_name]} rows;  '
      f'{b_name} {tot[b_name]:.1f}us/{cnt[b_name]} rows')
